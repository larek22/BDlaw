"""DOCX importer using python-docx."""

from __future__ import annotations

from pathlib import Path
from typing import Dict, Iterable

import docx

from bdlaw.ingest.base import Document, Importer


class DOCXImporter(Importer):
    mime_types: Iterable[str] = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    extensions: Iterable[str] = (".docx",)

    def supports(self, path: Path, mime_type: str) -> bool:
        return path.suffix.lower() in self.extensions or mime_type in self.mime_types

    def extract(self, path: Path) -> Document:
        doc = docx.Document(path)
        paragraphs = [p.text for p in doc.paragraphs]
        raw_text = "\n".join(paragraphs)
        metadata: Dict[str, object] = {"paragraph_count": len(paragraphs)}
        return Document(
            path=path,
            mime_type=self.mime_types[0],
            raw_text=raw_text,
            metadata=metadata,
        )


__all__ = ["DOCXImporter"]
